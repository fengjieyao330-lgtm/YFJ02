"""cv_create.py

在指定路径创建与 Jason_CV_2026_v9.docx 内容和格式完全一致的 Word 版简历。

用法:
    python cv_create.py <输出路径.docx> [照片路径.png]

默认输出: 脚本同目录下 Jason_CV_2026_v9.docx
默认照片: E:\\users\\YaoFJ01.CATLBATTERY\\Pictures\\IMG_2240.PNG
"""

import os
import sys

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH

DEFAULT_OUTPUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Jason_CV_2026_v9.docx')
DEFAULT_PHOTO = r'E:\users\YaoFJ01.CATLBATTERY\Pictures\IMG_2240.PNG'

USABLE_WIDTH_CM = 17.4

HEADER_LEFT = (
    '姚冯杰',
    '电话：+86 18692250432   ｜   邮箱：fengjieyao330@gmail.com',
    '出生年月：2002.10   ｜   籍贯：浙江绍兴',
)

SECTIONS = [
    ('教育经历', [
        ('entry', '2024.09-2025.11', '香港大学（HKU）电子与电气工程学院', '电气工程（硕士）'),
        ('bullet', 'GPA', '2.80；主修课程：机器学习、智慧电网、可再生能源科学、照明灯原理、投资学'),
        ('entry', '2020.09-2024.06', '中南大学（CSU）能源科学与工程学院', '能源与动力工程（本科）'),
        ('bullet', 'GPA', '83.01/100；主修课程：电子技术、机械设计、C语言、工程流体力学、热力学'),
    ]),
    ('工作/实习经历', [
        ('entry', '2025.12-至今', '宁德时代新能源科技股份有限公司', '工艺算法工程师（2025.10 实习入职，12 月转正）'),
        ('bullet', '质量智能体（2026.06-至今｜Agent 开发）', '开发 Neo4j FTA 知识图谱读写（670 节点）、Milvus 向量检索（3000+ chunks、768 维）等 MCP 工具与相关性分析矩阵，并从零开发 13 个 Agent 工具 REST 接口（图谱多跳遍历/向量+哈希双模检索）；构建结果展示闭环（自动落库、权限共享）'),
        ('bullet', '智能 FA 系统（2026.03-至今｜后端开发，数百人使用）', '① 独立完成 AI 看板（大模型对话式数据分析）：Java+Python 分工、异步队列+多轮会话，LLM 直接产出 ECharts 图表（存 S3），9 个 REST 接口，504 超时治理。② 30+ 业务 REST 接口开发与优化（全岗位累计独立开发 60+ REST 接口）：冷压断带智能监测与根因探索（异常卷判定、恶化 Top、原因与位置下钻、制程来料判责）、FTA 因子推理（失效特征回溯根因、部门级数据权限）、风险监测、改善经验库、工厂产品线、桌面云外网示警等；持续 SQL 调优、导出异步化（导出成功率 99%+）。③ 量产改善定时任务优化：新增第三步“根因推荐（原因+改善措施）暂存”，责任人确认后正式落库，单条失败不阻断整批'),
        ('bullet', '智能工艺设计（2026.03-2026.05｜后端开发）', '工艺文档向量化管理（Milvus/HNSW、文件级查重、启停）'),
        ('bullet', '厂房用电能耗项目（2025.12-2026.03）', '实习期参与需求调研与设备编码、工序量测等多源数据梳理；负责 Hive/Iceberg/Doris 数据治理（统一指标口径/数据质量）、构建 SQL+Airflow 数据流；参与 Python 数据平台产品化'),
        ('bullet', 'AI 提效', '熟练运用自研 OpenCode（接入 DeepSeek/GLM/Qwen/MiniMax 等）与 GitHub Copilot 辅助 SQL/Python/Java 开发'),
        ('entry', '2025.06-2025.07', '图迹科技有限公司', '新能源解决方案实习生'),
    ]),
    ('科研经历', [
        ('entry', '2025.01-2025.08', '基于需求响应的能源市场机制设计', '硕士毕设 · 独立完成'),
        ('bullet', '', 'Python+Gurobi 构建 TOU/RTP/动态定价电力市场-HEMS 仿真，动态定价机制：系统波动 -36.6%、可再生利用率 76.8%、电费 -18.4%'),
        ('entry', '2024.06', '三颗粒分离过程环索状液桥数值模拟', '本科毕设 · 院级优秀论文'),
        ('bullet', '', 'SurfaceEvolver 模拟三颗粒分离液桥演变（粒度比/液相体积/接触角 → 液桥力与断裂距离）'),
        ('entry', '2023-2024', '两篇独立作者（唯一作者）国际会议论文', '1 篇已收录 · 1 篇 EI 已提交'),
        ('bullet', '', '① AETS 2024：前后轴独立驱动车辆扭矩分配优化（PID+PSO）；② CONF-CIAP 2023：航空发动机分析'),
        ('entry', '2022.10-2023.05', '基于析氧电催化剂的锂电池正极材料高值化回收', '国家级大创 · 成员'),
        ('bullet', '', '去锂渣 SEM/EDS/XPS/XRD 表征，确定最优配比 SP:PVDF:去锂渣=1:1:9'),
    ]),
    ('专业技能', [
        ('bullet', '开发', 'Python、Java (Spring Boot)、SQL、FastAPI、MCP、Langfuse、Airflow、Postman、CAD、SurfaceEvolver、Origin'),
        ('bullet', 'AI 与数据', 'RAG、LLM 应用开发（提示词/Agent 工具/MCP）、Milvus、Neo4j、Redis、S3、Doris/Hive/Iceberg'),
        ('bullet', '语言', 'IELTS 7.5（流利商务沟通）、CET-6 504'),
    ]),
]


def setRunFont(run, size_pt, bold, font_cn):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = 'Times New Roman'
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), font_cn)
    return run


def addTextRun(paragraph, text, size_pt=10, bold=False, font_cn='宋体'):
    return setRunFont(paragraph.add_run(text), size_pt, bold, font_cn)


def setupNormalStyle(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(10)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.08


def setupSection(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.9)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)


def setCellMarginsZero(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for tag in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + tag)
        el.set(qn('w:w'), '0')
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    tcPr.append(margins)


def setCellWidth(cell, width_cm):
    cell.width = Cm(width_cm)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(int(width_cm * 567)))
    tcW.set(qn('w:type'), 'dxa')


def addHeaderTable(doc, photo_path):
    table = doc.add_table(rows=1, cols=2)
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tbl_pr.append(layout)
    left_cell, right_cell = table.rows[0].cells
    setCellWidth(left_cell, 13.8)
    setCellWidth(right_cell, 3.6)
    setCellMarginsZero(left_cell)
    setCellMarginsZero(right_cell)

    name_p = left_cell.paragraphs[0]
    name_p.paragraph_format.space_before = Pt(0)
    name_p.paragraph_format.space_after = Pt(3)
    addTextRun(name_p, HEADER_LEFT[0], 19, True, '黑体')

    phone_p = left_cell.add_paragraph()
    phone_p.paragraph_format.space_before = Pt(0)
    phone_p.paragraph_format.space_after = Pt(0)
    addTextRun(phone_p, HEADER_LEFT[1])

    birth_p = left_cell.add_paragraph()
    birth_p.paragraph_format.space_before = Pt(0)
    birth_p.paragraph_format.space_after = Pt(6)
    addTextRun(birth_p, HEADER_LEFT[2])

    photo_p = right_cell.paragraphs[0]
    photo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    photo_p.paragraph_format.space_before = Pt(2)
    photo_p.add_run().add_picture(photo_path, height=Cm(3.0))

    return table


def addSectionTitle(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    addTextRun(p, title, 12, True, '黑体')
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), '000000')
    pBdr.append(bottom)
    pPr.append(pBdr)


def addEntryLine(doc, date_text, org_text, right_text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(3)
    pf.space_after = Pt(0)
    pf.tab_stops.add_tab_stop(Cm(3.0), WD_TAB_ALIGNMENT.LEFT)
    pf.tab_stops.add_tab_stop(Cm(USABLE_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    addTextRun(p, date_text + '\t')
    addTextRun(p, org_text + '\t', bold=True)
    addTextRun(p, right_text, bold=True)


def addBulletItem(doc, lead_text, detail_text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.55)
    pf.first_line_indent = Cm(-0.33)
    pf.space_after = Pt(1.2)
    pf.tab_stops.add_tab_stop(Cm(0.55), WD_TAB_ALIGNMENT.LEFT)
    addTextRun(p, '•\t')
    if lead_text:
        addTextRun(p, lead_text + '：', bold=True)
    addTextRun(p, detail_text)


def createDocument(output_path, photo_path=DEFAULT_PHOTO):
    if not os.path.isfile(photo_path):
        raise FileNotFoundError('照片文件不存在: %s' % photo_path)
    doc = Document()
    setupNormalStyle(doc)
    setupSection(doc)
    addHeaderTable(doc, photo_path)
    for title, blocks in SECTIONS:
        addSectionTitle(doc, title)
        for block in blocks:
            if block[0] == 'entry':
                _, date_text, org_text, right_text = block
                addEntryLine(doc, date_text, org_text, right_text)
            else:
                _, lead_text, detail_text = block
                addBulletItem(doc, lead_text, detail_text)
    doc.save(output_path)


def main():
    output_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_OUTPUT
    photo_path = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_PHOTO
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    createDocument(output_path, photo_path)
    print('CV 已生成: %s' % output_path)


if __name__ == '__main__':
    main()

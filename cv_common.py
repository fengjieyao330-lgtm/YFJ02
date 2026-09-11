"""cv_common.py

各方向简历脚本共享的排版/工具函数与默认值。

- SECTIONS 结构约定：
  [ (标题, [ block, ... ]), ... ]
  block = ('entry', 日期, 机构, 右侧文字)
        | ('bullet', 引导词, 详情文字)     # 引导词为空时不显示加粗前缀
  entry 行 = 左：日期 + 中：机构（加粗） + 右：右侧文字（加粗），tab 对齐
"""

import os
import sys
import time
from io import BytesIO

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH

DEFAULT_PHOTO = r'E:\users\YaoFJ01.CATLBATTERY\Pictures\IMG_2240.PNG'

USABLE_WIDTH_CM = 17.4


# ---------------------------------------------------------------- 基础排版

def setRunFont(run, size_pt, bold, font_cn):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = 'Times New Roman'
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), font_cn)
    return run


def addTextRun(paragraph, text, size_pt=10, bold=False, font_cn='宋体'):
    return setRunFont(paragraph.add_run(text), size_pt, bold, font_cn)


def setupNormalStyle(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(10)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0


def setupSection(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.8)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)


def setCellMarginsZero(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for tag in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + tag)
        el.set(qn('w:w'), '0')
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    tcPr.append(margins)


def setCellWidth(cell, width_cm):
    cell.width = Cm(width_cm)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(int(width_cm * 567)))
    tcW.set(qn('w:type'), 'dxa')


# ---------------------------------------------------------------- 文件读取

def readPhotoStream(photo_path, retries=5, delay=2):
    """读取照片并校验文件头，瞬时占用/读取损坏时自动重试，返回 BytesIO"""
    last_err = None
    for i in range(retries):
        try:
            with open(photo_path, 'rb') as f:
                data = f.read()
        except Exception as e:
            last_err = e
        else:
            if data[:8] == b'\x89PNG\r\n\x1a\n' or data[:3] == b'\xff\xd8\xff':
                return BytesIO(data)
            last_err = ValueError('照片读取内容头部无效')
        time.sleep(delay)
    raise OSError('照片文件多次重试仍读取失败（可能被临时占用）: %s' % photo_path) from last_err


def isPathLocked(path):
    """检测文件是否被其他进程占用（能打开即未占用）；文件不存在返回 False"""
    if not os.path.exists(path):
        return False
    try:
        with open(path, 'a+b'):
            return False
    except OSError:
        return True


# ---------------------------------------------------------------- 文档组件

def addHeaderTable(doc, header_left, photo_path):
    """header_left: (姓名, 联系方式行, 第三行)"""
    table = doc.add_table(rows=1, cols=2)
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tbl_pr.append(layout)
    left_cell, right_cell = table.rows[0].cells
    setCellWidth(left_cell, 13.8)
    setCellWidth(right_cell, 3.6)
    setCellMarginsZero(left_cell)
    setCellMarginsZero(right_cell)

    name_p = left_cell.paragraphs[0]
    name_p.paragraph_format.space_before = Pt(0)
    name_p.paragraph_format.space_after = Pt(3)
    addTextRun(name_p, header_left[0], 19, True, '黑体')

    phone_p = left_cell.add_paragraph()
    phone_p.paragraph_format.space_before = Pt(0)
    phone_p.paragraph_format.space_after = Pt(0)
    addTextRun(phone_p, header_left[1])

    birth_p = left_cell.add_paragraph()
    birth_p.paragraph_format.space_before = Pt(0)
    birth_p.paragraph_format.space_after = Pt(6)
    addTextRun(birth_p, header_left[2])

    photo_p = right_cell.paragraphs[0]
    photo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    photo_p.paragraph_format.space_before = Pt(2)
    photo_p.add_run().add_picture(readPhotoStream(photo_path), height=Cm(3.0))

    return table


def addSectionTitle(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    addTextRun(p, title, 12, True, '黑体')
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), '000000')
    pBdr.append(bottom)
    pPr.append(pBdr)


def addEntryLine(doc, date_text, org_text, right_text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(0)
    pf.tab_stops.add_tab_stop(Cm(USABLE_WIDTH_CM / 2), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(USABLE_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    addTextRun(p, date_text + '\t')
    addTextRun(p, org_text + '\t', bold=True)
    addTextRun(p, right_text, bold=True)


def addBulletMarker(paragraph):
    """圆点 run + 精确 5.85pt 字符间距，替代 tab（规避 Word 个别内容下 tab 塌陷的排版怪癖）"""
    run = addTextRun(paragraph, '•')
    rpr = run._element.get_or_add_rPr()
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:val'), '117')
    sz = rpr.find(qn('w:sz'))
    if sz is not None:
        sz.addprevious(spacing)
    else:
        rpr.append(spacing)
    return run


def addBulletItem(doc, lead_text, detail_text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.55)
    pf.first_line_indent = Cm(-0.33)
    pf.space_after = Pt(1)
    pPr = p._p.get_or_add_pPr()
    op = OxmlElement('w:overflowPunct')
    op.set(qn('w:val'), '0')  # 禁止标点悬挂出右边距，保持各行右缘整齐
    anchor = pPr.find(qn('w:spacing'))
    if anchor is None:
        anchor = pPr.find(qn('w:ind'))
    if anchor is not None:
        anchor.addprevious(op)
    else:
        pPr.append(op)
    addBulletMarker(p)
    if lead_text:
        addTextRun(p, lead_text + '：', bold=True)
    addTextRun(p, detail_text)


# ---------------------------------------------------------------- 组装/输出

def createDocument(output_path, header_left, sections, photo_path=DEFAULT_PHOTO):
    if not os.path.isfile(photo_path):
        raise FileNotFoundError('照片文件不存在: %s' % photo_path)
    doc = Document()
    setupNormalStyle(doc)
    setupSection(doc)
    addHeaderTable(doc, header_left, photo_path)
    for title, blocks in sections:
        addSectionTitle(doc, title)
        for block in blocks:
            if block[0] == 'entry':
                _, date_text, org_text, right_text = block
                addEntryLine(doc, date_text, org_text, right_text)
            else:
                _, lead_text, detail_text = block
                addBulletItem(doc, lead_text, detail_text)
    doc.save(output_path)


def docxToPdf(docx_path, pdf_path=None, retries=3, delay=1.5):
    """用本机 Word 将 docx 另存为 PDF；pdf_path 缺省为与 docx 同名。

    目标文件被占用时明确提示并跳过（不崩溃）；返回 PDF 路径或 None
    """
    if pdf_path is None:
        pdf_path = os.path.splitext(os.path.abspath(docx_path))[0] + '.pdf'
    pdf_path = os.path.abspath(pdf_path)  # Word COM 的 SaveAs 需要绝对路径
    if isPathLocked(pdf_path):
        print('警告: PDF 正被其他程序占用，无法另存（请关闭正在查看该 PDF 的'
              '浏览器标签页、资源管理器预览窗格等，再重新运行）: %s' % pdf_path)
        return None
    if os.path.exists(pdf_path):
        try:
            os.remove(pdf_path)
        except OSError as e:
            print('警告: 旧 PDF 删除失败（文件被其他程序占用，关闭占用程序后重试）: %s' % e)
            return None
    try:
        import win32com.client
        import pywintypes
    except ImportError:
        print('警告: 未安装 pywin32，跳过 PDF 生成（pip install pywin32）')
        return None
    word = None
    try:
        word = win32com.client.Dispatch('Word.Application')
        word.Visible = False
        word.DisplayAlerts = 0  # wdAlertsNone：避免弹窗导致"命令失败"
        doc = word.Documents.Open(os.path.abspath(docx_path), ReadOnly=True)
        last_err = None
        for _ in range(retries):
            try:
                doc.SaveAs(pdf_path, FileFormat=17)  # wdFormatPDF = 17
                doc.Close(False)
                return pdf_path
            except (pywintypes.com_error, OSError) as e:
                last_err = e
                time.sleep(delay)
        print('警告: PDF 另存失败（常见原因：目标文件被其他程序占用，'
              '关闭占用程序后重试即可；如被企业安全策略拦截，'
              '请勿尝试改名或换目录绕过）: %s' % last_err)
        try:
            doc.Close(False)
        except Exception:
            pass
        return None
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass


def runGenerate(default_output, header_left, sections, photo_path=DEFAULT_PHOTO):
    """通用 CLI 入口：python <脚本> [输出.docx] [照片.png] [输出.pdf]"""
    output_path = sys.argv[1] if len(sys.argv) > 1 else default_output
    photo = sys.argv[2] if len(sys.argv) > 2 else photo_path
    pdf_path = sys.argv[3] if len(sys.argv) > 3 else None
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    if isPathLocked(output_path):
        print('错误: 目标 docx 正被其他程序占用（请关闭 Word 中打开的该文件后重试）: %s' % output_path)
        return
    createDocument(output_path, header_left, sections, photo)
    print('CV 已生成: %s' % output_path)
    pdf_out = docxToPdf(output_path, pdf_path)
    if pdf_out:
        print('PDF 已生成: %s' % pdf_out)
